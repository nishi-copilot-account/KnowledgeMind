"""
DOCX document reader.
"""

from pathlib import Path

from docx import Document

from app.ingestion.readers.base_reader import BaseReader
from app.knowledge.models import (
    KnowledgeDocument,
    KnowledgeMetadata,
    KnowledgePage,
)
from app.utils.logger import logger


class DOCXReader(BaseReader):
    """
    Reads Microsoft Word (.docx) documents.
    """

    def read(self, file_path: str, original_filename: str | None = None) -> KnowledgeDocument:

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"{path} does not exist.")

        logger.info("DOCX size: %s bytes", path.stat().st_size)
        logger.info("Opening DOCX : %s", path.name)

        try:

            document = Document(path)

        except Exception as ex:

            logger.exception(
                "Unable to open DOCX"
            )

            raise RuntimeError(
                f"{path.name} is not a valid DOCX document."
            ) from ex

        paragraphs = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if not text:
                continue

            if paragraph.style.name.startswith("Heading"):
                paragraphs.append(f"\n{text}\n")
            else:
                paragraphs.append(text)

        page_text = "\n".join(paragraphs)

        logger.info(
            "Extracted %s characters from DOCX",
            len(page_text),
        )

        core = document.core_properties

        metadata = KnowledgeMetadata(
            title=core.title or path.stem,
            author=core.author or "",
            subject=core.subject or "",
            keywords=core.keywords or "",
            page_count=1,
      )

        page = KnowledgePage(
            page_number=1,
            text=page_text,
        )

        logger.info("DOCX Loaded Successfully")

        # TODO:
        # Extract embedded images from DOCX.
        # Reuse ImageProcessor + OCRProcessor.

        # TODO:
        # Extract Word tables into KnowledgeTable objects.

        return KnowledgeDocument(
            filename=original_filename or path.name,
            metadata=metadata,
            pages=[page],
        )